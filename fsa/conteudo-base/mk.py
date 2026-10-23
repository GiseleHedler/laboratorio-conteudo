from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
wb=Workbook()
H=PatternFill('solid',fgColor='241A21'); HF=Font(bold=True,color='FFFFFF'); T=Font(bold=True,size=16,color='C2185B'); B=Font(bold=True)
IN=PatternFill('solid',fgColor='FFF4E5'); FO=PatternFill('solid',fgColor='EEF4FB')
th=Side(style='thin',color='CCCCCC'); BD=Border(top=th,bottom=th,left=th,right=th)
W=Alignment(wrap_text=True,vertical='top')
def head(ws,r,cols,widths=None):
    for i,c in enumerate(cols,1):
        x=ws.cell(r,i,c); x.fill=H; x.font=HF; x.alignment=Alignment(wrap_text=True,vertical='center'); x.border=BD
    if widths:
        for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i) if i<27 else 'A'+chr(64+i-26)].width=w
def col(i): 
    s='';
    while i: i,r=divmod(i-1,26); s=chr(65+r)+s
    return s

# 1 COMO FUNCIONA
ws=wb.active; ws.title='1. Como funciona'
ws.column_dimensions['A'].width=28; ws.column_dimensions['B'].width=95
ws['A1']='Conteúdo e Base FSA · Como funciona'; ws['A1'].font=T
rows=[
('AS CAMADAS (ler primeiro)','O PRINCIPAL DA SEMANA é transformar seguidor em base (e-mail). O número que mais importa: e-mails novos + cliques na newsletter.'),
('CAMADA 1 · O CENTRO','Tema da semana (vem da live) → NEWSLETTER. É o produto da semana e o que mantém a base quente.'),
('CAMADA 2 · A PONTE','Criativo de captura ("comenta e recebe"). Leva a pessoa do Instagram para dentro da base.'),
('CAMADA 3 · O COMBUSTÍVEL','Reels, carrosséis e YouTube. Atraem gente nova e aquecem quem já segue. Uns 3 por semana no tema da semana; o resto nos assuntos fixos ligados às formações (para o perfil não ficar repetitivo).'),
('CARROSSEL DA NEWSLETTER','NUNCA abrir com "o tema da newsletter desta semana". A capa é a ideia mais surpreendente que está DENTRO da newsletter, algo que quebra a cabeça e faz rolar. Só no último slide a pessoa descobre que aquilo é um pedaço da newsletter: "Isso saiu na nossa newsletter desta semana. Quer receber as próximas? Comenta NEWS."'),
('COMO SOMOS MEDIDOS (3 níveis)','1) TODA SEMANA, ENTREGAS (aba 3): fiz o que combinei? Meta 100%. 2) TODA SEMANA, KPIs (aba 5): os números que mostram se a operação está saudável (seguidores novos, engajados, e-mails novos, cliques). Eu trago e comparo com a semana anterior. 3) TODO MÊS, OKR (aba 2): a meta de RESULTADO do mês. É o que diz se o trabalho trouxe resultado. E em cada peça, o PLACAR (aba 4): gancho, retenção e conversão de cada reel e carrossel.'),
('OKR x KPI','OKR é o que queremos MUDAR no mês (crescer a base, aumentar cliques). KPI é o que não podemos deixar PIORAR toda semana. OKR muda a cada mês; KPI é sempre o mesmo painel.'),
('VISÃO','Reaquecer a base da FSA e deixá-la pronta para a próxima oferta. Instagram atrai e aquece; e-mail (newsletter) é a base que é nossa.'),
('RESPONSÁVEL','Uma pessoa: roteiros de reels, carrosséis, newsletter e criativo de captura. Ajuda na copy quando a Gisele pedir.'),
('QUEM APARECE','Reels e carrosséis divididos entre o Gabriel de Carvalho e a professora Liliane. Ela distribui entre os dois.'),
('O CENTRO DA SEMANA','A live semanal (aberta para toda a base). O tema da live vira a newsletter da semana. NÃO usar cortes da live como reels.'),
('ENTREGAS DA SEMANA','5 roteiros de reels + 4 carrosséis + 1 newsletter + YouTube (5 Shorts + 1 vídeo roteirizado + a live da semana). Um dos reels ou carrosséis é o criativo de captura ("comenta X e recebe Y").'),
('YOUTUBE: SHORTS','Os 5 reels da semana sobem também como Shorts no YouTube. Sem roteiro novo: é repostar, com título próprio para o YouTube.'),
('YOUTUBE: A LIVE DA SEMANA','A gravação da live sobe no YouTube toda semana. O trabalho dela: título forte (com a palavra que a pessoa buscaria no Google), capa com foto real e texto grande, descrição com o resumo, os momentos principais e o link da newsletter. Tempo: uns 30 minutos.'),
('YOUTUBE: VÍDEO ROTEIRIZADO','1 por semana, de 8 a 15 minutos, gravado pelo Gabriel ou pela Liliane. Usa o MESMO tema da newsletter da semana, mais aprofundado: a pesquisa já está feita, então o trabalho é o roteiro. Estrutura: gancho nos primeiros 30 segundos (a promessa do vídeo), desenvolvimento em blocos, caso prático, chamada para a newsletter.'),
('YOUTUBE: THUMB (CAPA)','Toda thumb tem ROSTO, nunca só nome ou texto. Rosto grande, com expressão que chama atenção (surpresa, dúvida, preocupação, apontando). Foto real do professor, feita para a thumb. Até 4 palavras em letra grande, que completam o título sem repetir. Flecha ou círculo apontando para o elemento principal (o exame, o alimento, o número). Contraste forte. Testar se dá para ler no tamanho de celular.'),
('YOUTUBE: O PAPEL','Quem busca no YouTube quer conteúdo longo e pesquisa por assunto. Todo vídeo leva para a newsletter: link na descrição e no comentário fixado. O número que mostra se a thumb funciona é a TAXA DE CLIQUES (CTR) das impressões, no YouTube Studio.'),
('CRIATIVO DE CAPTURA (rodízio)','Semana 1: fica na DM (engajamento). Semana 2: DM pede o e-mail e o material chega por e-mail (base). Semana 3: carrossel com o tema da newsletter + convite "comenta NEWS". Semana 4: DM para e-mail com material de outra formação.'),
('MATERIAL DE CAPTURA','Sempre ligado a uma formação (exames, suplementação, hormonal, fertilidade, jejum, alergias). Nada genérico. Conteúdo técnico tirado do material da FSA (live ou aula), nunca de fonte de fora sem checar.'),
('ROTINA (D = dia da live)','D+1: transcrição (Descript) e rascunho da newsletter. SEXTA-FEIRA: manda numa mensagem só o PACOTE DA PRÓXIMA SEMANA (roteiros, carrosséis, newsletter, captura) + o RELATÓRIO DA SEMANA (aba 5 preenchida). Segunda a quinta: produz, publica e preenche o placar (aba 4). Newsletter: escrita a partir da live, enviada até 2 dias depois dela.'),
('IMAGENS (responsabilidade dela)','Todo roteiro e carrossel sai com as imagens escolhidas, não só o texto. Imagem REAL, nunca de IA: professores, aulas, bastidores, prova (depoimentos e prints de alunos), e imagem daquilo que se está falando (o exame, o suplemento, o alimento). Muita imagem.'),
('BANCO DE IMAGENS','Usar primeiro os bancos próprios da FSA (fotos e vídeos de professores, aulas, bastidores, prova). O Elements só para o que é genérico (alimento, laboratório, fundo).'),
('ARTE DO CARROSSEL','A Claude monta o carrossel (skill carrossel FSA) a partir do texto e das fotos que ELA envia. Nenhuma imagem é gerada por IA: as fotos vêm sempre dos bancos da FSA ou do Elements.'),
('AUTORIZAÇÃO DE IMAGEM','Depoimento e print de aluno só com autorização. Exame de paciente só sem nome, data de nascimento ou qualquer dado que identifique.'),
('APROVAÇÃO','Tudo passa pela Gisele antes de sair, em LOTE, uma vez por semana: na SEXTA, o pacote da semana seguinte. A primeira entrega é na segunda-feira, 26/10.'),
('PRIORIDADE SE O TEMPO APERTAR','1) Newsletter  2) Criativo de captura  3) Vídeo roteirizado do YouTube  4) Carrosséis  5) Reels  6) Live no YouTube  7) Shorts'),
('A META É RESULTADO','Toda peça tem placar (aba 4). Gancho, retenção e conversão são medidos em cada reel. Não basta entregar: tem que superar a própria média.'),
('TODA SEMANA, 3 LINHAS','O que funcionou. O que não funcionou. O que vou testar na semana que vem. (aba 5)'),
('REGRAS QUE NÃO MUDAM','Mesma mensagem da página de destino em toda peça de campanha. Fatos com fonte. Sem promessa de cura ou resultado. Respeitar as regras do conselho. Sem travessão. Português do Brasil.'),
('BASE DE E-MAIL','Reaquecer em ordem: semanas 1 e 2 só para quem abriu ou clicou nos últimos 6 meses; semana 3 campanha de reativação para os parados; depois tirar do envio quem não abre há 12 meses. Spam acima de 0,1%: parar e corrigir.'),
]
for i,(a,b) in enumerate(rows,3):
    ws.cell(i,1,a).font=B; ws.cell(i,2,b).alignment=W; ws.cell(i,1).alignment=W
    ws.row_dimensions[i].height=max(30,15*(len(b)//90+1))

# 2 OKR DO MES
ws=wb.create_sheet('2. OKR do mês')
ws['A1']='OKR do mês'; ws['A1'].font=T
ws['A2']='Mês:'; ws['B2'].fill=IN
ws['A3']='Objetivo:'; ws['B3']='Reaquecer a base da FSA e deixá-la pronta para a próxima oferta.'; ws['B3'].font=B
head(ws,5,['Resultado-chave','Como medir','Linha de base (mês anterior)','Meta do mês','Realizado','% atingido','Comentário'],[34,46,18,14,14,12,40])
krs=[('KR1. Gancho e retenção dos reels','% dos reels do mês com retenção (tempo médio ÷ duração) acima da linha de base'),
('KR2. Base nova vinda do Instagram','E-mails novos gerados pelos criativos de captura no mês'),
('KR3. Engajamento que importa','(Comentários + salvamentos + compartilhamentos) ÷ alcance, média do mês'),
('KR4. Newsletter viva','Taxa média de cliques das newsletters do mês'),
('KR5. Base reativada','Contatos parados que voltaram a abrir ou clicar no mês')]
for i,(a,b) in enumerate(krs,6):
    ws.cell(i,1,a).font=B; ws.cell(i,2,b).alignment=W
    for c in (3,4,5): ws.cell(i,c).fill=IN
    ws.cell(i,6,f'=IFERROR(E{i}/D{i},"")').number_format='0%'
    for c in range(1,8): ws.cell(i,c).border=BD
    ws.row_dimensions[i].height=32
ws['A12']='Indicador de apoio (não é meta):'; ws['A12'].font=B; ws['B12']='Seguidores novos no mês. Acompanhar, mas não perseguir: meta de seguidor puxa conteúdo que viraliza e não vende.'; ws['B12'].alignment=W
ws['A14']='Regra:'; ws['A14'].font=B; ws['B14']='No 1º mês, as duas primeiras semanas viram a linha de base. A partir do 2º mês, a meta é superar a linha de base do mês anterior. Bater 70% do OKR já é bom resultado.'; ws['B14'].alignment=W

# 3 ENTREGAS
ws=wb.create_sheet('3. Entregas da semana')
ws['A1']='Entregas da semana (meta: 100%, sob controle dela)'; ws['A1'].font=T
cols=['Semana (início)','Reels roteirizados (meta 5)','Carrosséis (meta 4)','Newsletter enviada (1)','Captura publicada (1)','Tipo de captura','% entregue','Pacote enviado p/ aprovação no prazo?','Observação','Shorts no YouTube (meta 5)','Live no YouTube (1)','Vídeo roteirizado no YouTube (1)']
head(ws,3,cols,[16,16,14,14,14,22,12,18,40,14,14,16])
dv=DataValidation(type='list',formula1='"Fica na DM,DM para e-mail,Convite newsletter"'); ws.add_data_validation(dv)
dv2=DataValidation(type='list',formula1='"Sim,Não"'); ws.add_data_validation(dv2)
for r in range(4,17):
    for c in range(1,13): ws.cell(r,c).border=BD
    for c in (1,2,3,4,5,6,8,9,10,11,12): ws.cell(r,c).fill=IN
    ws.cell(r,7,f'=IF(A{r}="","",(MIN(B{r},5)+MIN(C{r},4)+MIN(D{r},1)+MIN(E{r},1)+MIN(J{r},5)+MIN(K{r},1)+MIN(L{r},1))/18)').number_format='0%'
    ws.cell(r,7).fill=FO
    dv.add(f'F{r}'); dv2.add(f'H{r}')

# 4 PLACAR POR PECA
ws=wb.create_sheet('4. Placar por peça')
ws['A1']='Placar por peça: gancho, retenção e conversão'; ws['A1'].font=T
ws['A2']='Laranja: ela preenche (Insights do Instagram, 7 dias após publicar). Azul: calcula sozinho.'
cols=['Semana','Data','Tipo','Quem aparece','Tema','Gancho (o que é dito/escrito nos 3 primeiros segundos)','Link',
'Alcance','Taxa de pulos % (reels)','Duração (s)','Tempo médio assistido (s)','RETENÇÃO %','Comentários','Salvamentos','Compartilhamentos','Seguidores ganhos','CONVERSÃO % (engajamento ÷ alcance)',
'Imagens reais? (Sim/Não)','É captura?','E-mails gerados','CAPTURA % (e-mails ÷ comentários)','Acima da linha de base?','O que aprendi']
head(ws,4,cols,[10,11,11,13,22,40,18,10,10,9,10,11,10,10,12,10,13,11,9,10,12,12,34])
ws.row_dimensions[4].height=60
dvT=DataValidation(type='list',formula1='"Reel,Carrossel,Newsletter,Short,Live YouTube,Vídeo YouTube"'); dvP=DataValidation(type='list',formula1='"Gabriel,Liliane,Outro"'); dvC=DataValidation(type='list',formula1='"Sim,Não"')
for d in (dvT,dvP,dvC): ws.add_data_validation(d)
ws['X2']='Linha de base de retenção %:'; ws['Y2'].fill=IN; ws['Y2'].number_format='0%'
ws['X3']='Linha de base de conversão %:'; ws['Y3'].fill=IN; ws['Y3'].number_format='0.0%'
ws.column_dimensions['X'].width=28
for r in range(5,45):
    for c in range(1,24): ws.cell(r,c).border=BD; ws.cell(r,c).fill=IN
    ws.cell(r,12,f'=IFERROR(K{r}/J{r},"")').number_format='0%'
    ws.cell(r,17,f'=IFERROR((M{r}+N{r}+O{r})/H{r},"")').number_format='0.0%'
    ws.cell(r,21,f'=IF(S{r}="Sim",IFERROR(T{r}/M{r},""),"")').number_format='0%'
    ws.cell(r,22,f'=IF(OR(H{r}="",$Y$3=""),"",IF(AND(OR(C{r}<>"Reel",L{r}>=$Y$2),Q{r}>=$Y$3),"Sim","Não"))')
    for c in (12,17,21,22): ws.cell(r,c).fill=FO
    dvT.add(f'C{r}'); dvP.add(f'D{r}'); dvC.add(f'R{r}'); dvC.add(f'S{r}')
ws.freeze_panes='F5'

# 5 KPIs
ws=wb.create_sheet('5. KPIs da semana')
ws['A1']='Relatório da semana (ela traz toda semana)'; ws['A1'].font=T
cols=['Semana (início)','Seguidores (total)','Seguidores novos','Contas engajadas','Alcance','E-mails novos (captura)','Newsletter: abertura %','Newsletter: cliques %','Descadastros','Spam %','Reels acima da linha de base','YouTube: inscritos novos','YouTube: visualizações','YouTube: horas assistidas','YouTube: CTR da thumb %','YouTube: cliques no link da newsletter','O que funcionou','O que não funcionou','O que vou testar na semana que vem']
head(ws,3,cols,[14,12,12,12,12,13,12,12,11,9,14,12,12,12,12,14,34,34,34])
ws.row_dimensions[3].height=45
for r in range(4,17):
    for c in range(1,20): ws.cell(r,c).border=BD; ws.cell(r,c).fill=IN
    ws.cell(r,11,f"=IF(A{r}=\"\",\"\",COUNTIFS('4. Placar por peça'!A:A,A{r},'4. Placar por peça'!C:C,\"Reel\",'4. Placar por peça'!V:V,\"Sim\")&\" de \"&COUNTIFS('4. Placar por peça'!A:A,A{r},'4. Placar por peça'!C:C,\"Reel\"))"); ws.cell(r,11).fill=FO
    for c in (17,18,19): ws.cell(r,c).alignment=W
ws['A19']='Atenção: olhar mais os CLIQUES do que a abertura (o iPhone marca e-mail como aberto sozinho). Spam acima de 0,1%: parar e avisar a Gisele.'
ws['A20']='Na coluna "Semana" das abas 3, 4 e 5, usar sempre a mesma data (a segunda-feira da semana), para as contas cruzarem.'

# 6 CHECKLIST
ws=wb.create_sheet('6. Checklist do roteiro')
ws['A1']='Checklist antes de mandar o roteiro para aprovação'; ws['A1'].font=T
ws.column_dimensions['A'].width=6; ws.column_dimensions['B'].width=100
items=['GANCHO: escrevi 3 opções de abertura e escolhi a mais forte.',
'GANCHO: a primeira frase cabe em 3 segundos e cria tensão, curiosidade ou contradiz algo que a pessoa acredita.',
'GANCHO: tem texto grande na tela nos 3 primeiros segundos (muita gente assiste sem som).',
'IMAGEM: escolhi as imagens de cada parte (reais, sem IA) e elas mostram aquilo que está sendo dito.',
'IMAGEM: a primeira imagem/cena é a mais forte e já se entende sem som.',
'IMAGEM: depoimento ou exame com autorização e sem dados que identifiquem.',
'GANCHO: não começa com "Oi, gente", apresentação ou contexto.',
'RETENÇÃO: cada frase dá motivo para ouvir a próxima. Cortei toda frase que não faz falta.',
'RETENÇÃO: a promessa do gancho só é paga no final (ou perto dele).',
'RETENÇÃO: o reel tem até 45 segundos, salvo motivo forte.',
'RETENÇÃO: marquei onde entra mudança visual (corte, texto, imagem) a cada 3 a 5 segundos.',
'CONVERSÃO: tem UM pedido só no final (comentar, salvar, seguir, compartilhar ou ir para a newsletter).',
'CONVERSÃO: se é captura, a palavra-chave e o material estão claros e a automação está pronta.',
'FATOS: todo dado técnico vem do material da FSA (live, aula) ou tem fonte.',
'CARROSSEL DA NEWSLETTER: a capa NÃO diz "tema da newsletter". Abre com a ideia que quebra a cabeça; a newsletter só aparece no último slide.',
'REGRAS: sem promessa de cura ou resultado garantido; dentro das regras do conselho.',
'YOUTUBE (thumb): tem rosto grande com expressão forte, até 4 palavras, flecha ou círculo no elemento principal, e dá para ler no celular. Nunca só nome ou texto.',
'YOUTUBE (vídeo roteirizado): a promessa do vídeo aparece nos primeiros 30 segundos.',
'YOUTUBE (live): título com a palavra que a pessoa buscaria, capa com foto real e texto grande, descrição com resumo, momentos principais e link da newsletter, comentário fixado com o link.',
'CAMPANHA: se fala de oferta, diz exatamente o que a página de destino diz (mesma frase, número, prazo).']
for i,t in enumerate(items,3):
    ws.cell(i,1,'☐'); ws.cell(i,2,t).alignment=W

# 7 CAPTURA
ws=wb.create_sheet('7. Calendário de captura')
ws['A1']='Calendário de captura (1 por semana, em rodízio)'; ws['A1'].font=T
head(ws,3,['Semana','Tipo','Formação ligada','Palavra-chave','Material (o que a pessoa recebe)','Tirado de qual material?','Comentários','E-mails gerados'],[14,20,22,16,50,14,12,12])
tipos=['Fica na DM','DM para e-mail','Convite newsletter','DM para e-mail']
for r in range(4,12):
    ws.cell(r,2,tipos[(r-4)%4])
    for c in range(1,9): ws.cell(r,c).border=BD
    for c in (1,3,4,5,6,7,8): ws.cell(r,c).fill=IN
wb.save('okr-conteudo-fsa.xlsx')
