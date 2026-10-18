from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE

IMG = '/home/user/laboratorio-conteudo/fsa/black-vitalicia/'
SERIF, SANS, BLACK = 'Cormorant Garamond', 'Archivo', 'Archivo Black'
C = dict(bg='14020D', bg2='1d0616', roxo='2B0739', creme='FBF6EF', texto='F1E9EC', suave='D2C1CA',
         ouro='E0B64F', ouro2='E9C56A', rosa='DC447C', linha='3A1630', claro='F4EDE3', escuro='2A0F1F')
W, H = 13.333, 7.5

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BL = prs.slide_layouts[6]
rgb = lambda h: RGBColor.from_string(h)


def novo(fundo='bg'):
    s = prs.slides.add_slide(BL)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(C[fundo])
    return s


def caixa(s, x, y, w, h, cor=None, borda=None, tracejado=False, forma=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    if cor: r.fill.solid(); r.fill.fore_color.rgb = rgb(C[cor])
    else: r.fill.background()
    if borda:
        r.line.color.rgb = rgb(C[borda]); r.line.width = Pt(2)
        if tracejado: r.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else: r.line.fill.background()
    r.shadow.inherit = False
    return r


def txt(s, x, y, w, h, partes, al='l', anc='t', esp=1.0):
    """partes: lista de parágrafos; cada parágrafo é lista de (texto, fonte, tamanho, cor, italico, negrito) ou str."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anc]
    for i, par in enumerate(partes):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[al]
        p.line_spacing = esp
        for run in par:
            t, f, sz, cor, it, bd = (run + (False, False))[:6] if isinstance(run, tuple) else (run, SANS, 20, 'texto', False, False)
            r = p.add_run(); r.text = t
            r.font.name = f; r.font.size = Pt(sz); r.font.color.rgb = rgb(C[cor]); r.font.italic = it; r.font.bold = bd
    return tb


def img(s, nome, x, y, w=None, h=None):
    return s.shapes.add_picture(IMG + nome if '/' not in nome else nome, Inches(x), Inches(y),
                                Inches(w) if w else None, Inches(h) if h else None)


def capa_img(s, nome, x, y, w, h, esq=0.5):
    from PIL import Image
    iw, ih = Image.open(IMG + nome).size
    p = s.shapes.add_picture(IMG + nome, Inches(x), Inches(y), Inches(w), Inches(h))
    ar, br = iw / ih, w / h
    if ar > br:
        c = (1 - br / ar); p.crop_left = c * esq; p.crop_right = c * (1 - esq)
    else:
        c = (1 - ar / br) / 2; p.crop_top = c; p.crop_bottom = c
    return p


def logo(s, x=0.6, y=0.45, h=0.42):
    img(s, 'logo-fsa-escuro.png', x, y, h=h)


def eyebrow(s, t, y=0.55, x=None, al='l', cor='ouro'):
    txt(s, x if x is not None else 0.6, y, 12.1 if x is None else 12.1 - x, 0.4, [[(t.upper(), SANS, 14, cor, False, True)]], al=al)


def chat(s, perguntas, titulo='Me conta no chat'):
    eyebrow(s, '● ' + titulo, cor='rosa')
    y = 1.7
    for q in perguntas:
        txt(s, 0.9, y, 11.5, 1.2, [[(q, SERIF, 50, 'creme')]]); y += 1.45


def frase(s, l1, l2=None, sz=72, cor2='ouro2', y=2.2, al='c'):
    par = [[(l1, SERIF, sz, 'creme')]]
    if l2: par.append([(l2, SERIF, sz, cor2, True)])
    txt(s, 0.8, y, 11.7, 3.4, par, al=al, anc='m', esp=0.95)


def pendente(s, x, y, w, h, titulo, detalhe):
    caixa(s, x, y, w, h, cor='roxo', borda='rosa', tracejado=True)
    txt(s, x + 0.35, y + 0.3, w - 0.7, h - 0.6,
        [[('PARA O TIME: ' + titulo, SANS, 18, 'rosa', False, True)], [(detalhe, SANS, 15, 'suave')]], anc='m', esp=1.15)


def nota(s, t):
    s.notes_slide.notes_text_frame.text = t


# 1 CAPA
s = novo()
capa_img(s, 'cena.jpg', 7.0, 0, W - 7.0, H, esq=0.12)
logo(s)
txt(s, 0.6, 1.9, 6.4, 4, [[('Festa dos 27 anos da', SERIF, 40, 'creme')], [('Nutrição Funcional', SERIF, 56, 'ouro2', True)],
                          [(' ', SANS, 14, 'creme')], [('A grande revelação da Black da FSA', SANS, 20, 'texto')]], esp=1.0)
txt(s, 0.6, 5.9, 6.4, 0.8, [[('19.10  ·  20H27  ·  AO VIVO', SANS, 18, 'ouro', False, True)]])
nota(s, 'Capa. Fica na tela no início.')

# 2 PRÉ-LIVE
s = novo('bg2'); logo(s)
frase(s, 'Já vamos começar.', 'Separa uma folha em branco e uma caneta.', sz=54, y=1.6)
txt(s, 0.8, 5.6, 11.7, 0.8, [[('Chama quem atende com você para assistir junto.', SANS, 22, 'suave')]], al='c')
nota(s, 'Pré-live (20h15 a 20h27). Intercalar com os vídeos de depoimento (Daniela, Lysiane, Juliana, Maria).')

# 3 CHAT ABERTURA
s = novo(); chat(s, ['De onde tu é?', 'Qual a tua profissão?', 'Desde quando acompanha a Nutrição Funcional?'])
nota(s, 'Bloco 1. Ler nomes, cidades e profissões. Comemorar quem acompanha há mais tempo.')

# 4 NÃO É SÓ BLACK
s = novo(); frase(s, 'Hoje não é só Black.', 'Festa de professor é com aula.')
nota(s, 'Bloco 1. "Depois da aula eu vou contar uma coisa. Eu e a Gi estávamos maturando isso faz tempo."')

# 5 POSSO COMEÇAR
s = novo('bg2')
frase(s, 'Pega a folha e a caneta.', 'Posso começar?', sz=64, y=1.5)
txt(s, 0.8, 5.5, 11.7, 0.8, [[('Me dá o OK no chat', SANS, 24, 'rosa', False, True)]], al='c')
nota(s, 'Bloco 2. Antes: pedir para chamarem colegas nos grupos e na comunidade do Facebook; Instagram vem para o YouTube.')

# 6 3 LAUDOS
s = novo(); eyebrow(s, '● Me conta no chat', cor='rosa')
txt(s, 0.9, 1.4, 11.5, 2.4, [[('Pensa nos últimos 3 laudos "normais" que chegaram até ti.', SERIF, 40, 'creme')],
                             [('Quantos desses pacientes estavam realmente bem?', SERIF, 40, 'ouro2', True)]], esp=1.0)
for i, n in enumerate('0123'):
    caixa(s, 1.6 + i * 2.7, 4.4, 2.1, 2.1, borda='ouro', forma=MSO_SHAPE.OVAL)
    txt(s, 1.6 + i * 2.7, 4.4, 2.1, 2.1, [[(n, SERIF, 80, 'ouro2')]], al='c', anc='m')
nota(s, 'Bloco 3. Comentar que quase ninguém responde 3.')

# 7 NORMAL PARA QUEM
s = novo(); frase(s, 'Normal', 'para quem?', sz=120, cor2='rosa', y=1.9)
nota(s, 'Bordão da noite. Repetir sempre que mostrar um exame "normal".')

# 8 NARRATIVA CEGA
s = novo('bg2'); eyebrow(s, 'Isso tem nome')
txt(s, 0.6, 1.3, 12, 1.6, [[('Narrativa Cega', SERIF, 96, 'ouro2', True)]])
txt(s, 0.6, 3.6, 11, 2.5, [[('Olhar só o número do laudo, só o protocolo pronto, e ir na tentativa e erro.', SERIF, 40, 'creme')]], esp=1.0)

# 9 NINGUÉM NASCE
s = novo(); frase(s, 'Ninguém nasce com a Narrativa Cega.', 'O sistema ensinou.', sz=64)
nota(s, 'Bloco 4.1. Pedir para repetirem em voz alta.')

# 10 GABRIEL: FAIXA DE REFERÊNCIA
s = novo('bg2'); eyebrow(s, 'De onde vem')
txt(s, 0.6, 1.1, 12, 1.2, [[('Como nasce o "normal" do laudo', SERIF, 52, 'creme')]])
pendente(s, 0.6, 2.6, 12.1, 4.2, 'CONTEÚDO DO GABRIEL',
         'Como a faixa de referência do laboratório é construída e por que "normal" não quer dizer "saudável". '
         'Gráfico, números e a fonte de cada dado vêm do Gabriel. Podem ser 1 a 3 slides. Não usar dado sem fonte.')

# 11 CHAT FACULDADE
s = novo(); chat(s, ['Qual frase tu mais ouviu', 'na faculdade sobre exame?'])
txt(s, 0.9, 5.2, 11.5, 1, [[('"Está dentro da referência."', SERIF, 40, 'ouro2', True)]])
nota(s, 'Bloco 4.1. A frase em dourado só aparece depois das respostas (animação ou slide seguinte).')

# 12 SINTOMAS
s = novo('bg2'); eyebrow(s, 'Tu reconhece quando ela aparece no consultório')
sint = ['Exames dentro da referência, e o paciente ainda com queixas que ninguém explicou.',
        'O ferro que tu suplementa por meses, ajusta a dose, troca a forma, e não sobe.',
        'O protocolo que funciona numa paciente e não funciona na outra.',
        'Onze queixas, uma lista de suplementos e nenhuma ordem para começar.']
for i, t in enumerate(sint):
    y = 1.35 + i * 1.45
    caixa(s, 0.6, y, 0.06, 1.1, cor='rosa')
    txt(s, 0.95, y, 11.6, 1.1, [[(t, SERIF, 32, 'creme')]], anc='m', esp=1.0)
nota(s, 'Bloco 4.1. Ler em voz alta e perguntar: "Bateu? Me conta no chat."')

# 13 NARRATIVA É CEGA
s = novo(); frase(s, 'A narrativa é cega.', 'Tu não é.', sz=96, cor2='rosa')
nota(s, 'Bloco 4.2. "A culpa não é tua. Tu foi ensinado a olhar o quadradinho."')

# 14 ICEBERG
s = novo('bg2'); eyebrow(s, 'O que o laudo mostra')
caixa(s, 0, 2.55, W, 0.03, cor='suave')
tri = caixa(s, 5.4, 1.2, 2.5, 1.35, cor='creme', forma=MSO_SHAPE.ISOSCELES_TRIANGLE)
txt(s, 5.4, 1.85, 2.5, 0.6, [[('O LAUDO', SANS, 16, 'escuro', False, True)]], al='c')
base = caixa(s, 2.4, 2.6, 8.5, 4.5, cor='roxo', forma=MSO_SHAPE.TRAPEZOID); base.rotation = 180
pal = ['sono', 'intestino', 'alimentação', 'estresse', 'história', 'como vive', 'onde mora', 'o que sente']
for i, p in enumerate(pal):
    txt(s, 2.9 + (i % 4) * 1.9, 3.4 + (i // 4) * 1.4, 1.9, 0.8, [[(p, SERIF, 30, 'ouro2', True)]], al='c')
txt(s, 9.3, 1.3, 3.5, 1, [[('O que o paciente é', SANS, 16, 'suave')]], al='r')

# 15 GABRIEL: CASO CLÍNICO
s = novo(); eyebrow(s, 'Um caso')
pendente(s, 0.6, 1.2, 12.1, 5.6, 'CASO CLÍNICO DO GABRIEL',
         'Um caso real explicado pelo iceberg: o que o laudo mostrava e o que estava embaixo. '
         'Sem nome, foto ou qualquer dado que identifique o paciente.')

# 16 EXERCÍCIO
s = novo('bg2'); eyebrow(s, 'Exercício · A folha da paciente')
passos = [('1', 'No meio da folha, desenha um quadradinho pequeno. Esse é o laudo.'),
          ('2', 'Fecha os olhos e lembra daquela paciente que voltou com o exame normal e continuava mal.'),
          ('3', 'Em volta do quadradinho, escreve tudo o que ela te contou que não estava no laudo.')]
for i, (n, t) in enumerate(passos):
    y = 1.4 + i * 1.85
    txt(s, 0.6, y, 1.2, 1.5, [[(n, SERIF, 96, 'ouro2')]], anc='m')
    txt(s, 2.0, y, 10.6, 1.5, [[(t, SERIF, 36, 'creme')]], anc='m', esp=1.0)
nota(s, 'Bloco 4.3. Ler o texto guiado do roteiro com calma. Deixar 2 a 3 minutos de silêncio com música baixa no passo 3.')

# 17 FOLHA
s = novo('claro')
caixa(s, 3.9, 0.5, 5.5, 6.5, cor='creme', borda='suave')
caixa(s, 6.25, 3.35, 0.8, 0.8, borda='escuro')
txt(s, 6.25, 3.35, 0.8, 0.8, [[('laudo', SANS, 10, 'escuro')]], al='c', anc='m')
em_volta = [('dorme mal', 4.3, 1.0), ('intestino preso', 6.9, 1.2), ('cansaço', 4.5, 2.2), ('come no carro', 7.3, 2.4),
            ('separação', 4.2, 3.5), ('trabalho noturno', 7.3, 3.5), ('ansiedade', 4.6, 4.8), ('remédio contínuo', 7.0, 4.7),
            ('incha', 5.3, 5.9), ('sem fome de manhã', 6.9, 6.0)]
for t, x, y in em_volta:
    txt(s, x, y, 2.2, 0.5, [[(t, SERIF, 24, 'escuro', True)]])
txt(s, 9.8, 2.4, 3.2, 2.5, [[('O laudo é o quadradinho.', SERIF, 28, 'escuro')], [('A paciente é tudo isso.', SERIF, 28, 'rosa', True)]], esp=1.0)
nota(s, 'Bloco 4.3. Mostrar depois que eles escreverem. Exemplo ilustrativo, não é um caso real.')

# 18 CHAT FOLHA
s = novo(); chat(s, ['Escreve no chat uma coisa', 'que está na tua folha.'])
nota(s, 'O Gabriel escolhe 2 ou 3 respostas e raciocina ao vivo: por onde começaria a investigar.')

# 19 HIPÓCRATES
s = novo()
img(s, 'marco-hipocrates.jpg', 0.6, 1.2, h=5.1)
eyebrow(s, '2.400 anos atrás', x=6.3)
txt(s, 6.3, 1.4, 6.4, 4.5, [[('Hipócrates já perguntava ao paciente', SERIF, 40, 'creme')],
                            [('o que ele comia, como vivia, onde morava.', SERIF, 40, 'ouro2', True)],
                            [(' ', SANS, 12, 'creme')], [('A gente não inventou nada. A gente só voltou a olhar.', SANS, 20, 'suave')]], esp=1.0)

# 20 OU TU JÁ SABE
s = novo('bg2'); frase(s, 'Ou tu já sabe isso,', 'ou tu ainda não sabe nada.', sz=72)

# 21 EU TRATO GENTE
s = novo(); frase(s, 'Eu não trato exame.', 'Eu trato gente.', sz=100, cor2='rosa')
nota(s, 'Pedir para repetirem em voz alta.')

# 22 POSTA A FOLHA
s = novo('bg2'); frase(s, 'Tira uma foto da tua folha.', 'Posta nos stories e marca @fsa_saudeavancada', sz=54)

# 23 0 A 10
s = novo(); eyebrow(s, '● Me conta no chat', cor='rosa')
txt(s, 0.9, 1.4, 11.5, 2.6, [[('De 0 a 10, quanta segurança tu sente quando o exame vem normal', SERIF, 44, 'creme')],
                             [('e a paciente não melhora?', SERIF, 44, 'ouro2', True)]], esp=1.0)
for i in range(11):
    caixa(s, 0.75 + i * 1.08, 5.0, 0.9, 0.9, borda='ouro', forma=MSO_SHAPE.OVAL)
    txt(s, 0.75 + i * 1.08, 5.0, 0.9, 0.9, [[(str(i), SERIF, 32, 'ouro2')]], al='c', anc='m')

# 24 1999
s = novo('bg2')
pendente(s, 0.6, 0.8, 5.6, 5.9, 'FOTO DO GABRIEL EM 1999',
         'A foto de quando ele começou a ensinar Nutrição Funcional no Brasil. Vertical ou quadrada, a melhor resolução que tiverem.')
txt(s, 6.8, 1.5, 6, 2, [[('1999', SERIF, 140, 'ouro2')]])
txt(s, 6.8, 3.9, 6, 2.5, [[('O começo da Nutrição Funcional no Brasil.', SERIF, 36, 'creme')]], esp=1.0)
nota(s, 'Bloco 5. O Gabriel conta a história verdadeira de 1999 (1 a 2 minutos).')

# 25 NÚMEROS
s = novo(); logo(s)
nums = [('27 anos', 'ensinando Nutrição Funcional'), ('40 mil+', 'alunos já estudaram na FSA'), ('1999', 'introduziu a Nutrição Funcional no Brasil')]
for i, (n, d) in enumerate(nums):
    x = 0.6 + i * 4.15
    txt(s, x, 2.2, 3.9, 1.6, [[(n, SERIF, 80, 'ouro2')]])
    txt(s, x, 3.9, 3.7, 1.2, [[(d, SANS, 20, 'texto')]], esp=1.1)
txt(s, 0.6, 5.7, 12, 0.8, [[('Primeiro profissional da América Latina certificado pelo The Institute for Functional Medicine (EUA)', SANS, 16, 'suave')]])

# 26 NUNCA FIZ
s = novo('bg2'); img(s, 'gabriel-fsa.jpg', 8.9, 0, h=H)
txt(s, 0.6, 1.6, 8, 4.2, [[('Em 27 anos ensinando,', SERIF, 56, 'creme')], [('eu nunca fiz o que', SERIF, 56, 'creme')],
                          [('vou fazer agora.', SERIF, 56, 'ouro2', True)]], anc='m', esp=0.95)

# 27 CHAT ACESSO
s = novo(); eyebrow(s, '● Me conta no chat', cor='rosa')
txt(s, 0.9, 1.6, 11.5, 4, [[('O acesso de alguma formação já venceu antes de tu terminar?', SERIF, 48, 'creme')],
                           [('Ou tu quis voltar numa aula e não dava mais?', SERIF, 48, 'ouro2', True)]], esp=1.0)

# 28 HOJE 2 ANOS
s = novo('bg2'); eyebrow(s, 'Hoje')
txt(s, 0.6, 1.8, 12, 2, [[('2 anos', SERIF, 150, 'suave')]])
txt(s, 0.6, 4.4, 12, 1.5, [[('de acesso em cada formação da FSA.', SERIF, 44, 'creme')]])

# 29 VITALÍCIO
s = novo('bg'); eyebrow(s, 'Nesta Black, pela primeira vez', y=1.6, al='c', cor='ouro')
txt(s, 0.3, 2.3, 12.7, 2.8, [[('VITALÍCIO', BLACK, 120, 'creme')]], al='c', anc='m')
txt(s, 0.3, 5.3, 12.7, 0.8, [[('As formações ficam contigo. Para sempre.', SERIF, 36, 'ouro2', True)]], al='c')
nota(s, 'Bloco 6. A revelação. Pausa antes e depois.')

# 30 CASO DIFÍCIL
s = novo('bg2'); frase(s, 'O caso difícil não avisa quando chega.', 'Quando ele chegar, a aula vai estar lá.', sz=56)

# 31 NÃO É AULA VELHA
s = novo(); eyebrow(s, 'Vitalício, do jeito certo')
txt(s, 0.6, 1.1, 12, 1.2, [[('Não é depósito de aula velha.', SERIF, 54, 'creme')]])
pts = [('A formação inteira', 'a mesma que a FSA vende hoje.'),
       ('As atualizações dela', 'aulas revisadas e conteúdos novos entram no teu acesso, sem custo.'),
       ('Sem letra miúda', 'formação nova, com outro nome ou outra versão (uma 4.0), é outro produto.')]
for i, (a, b) in enumerate(pts):
    y = 2.7 + i * 1.45
    caixa(s, 0.6, y, 0.06, 1.1, cor='ouro')
    txt(s, 0.95, y, 11.5, 1.1, [[(a, SANS, 24, 'ouro2', False, True)], [(b, SANS, 20, 'texto')]], anc='m', esp=1.1)

# 32 TRÊS PACOTES
s = novo(); eyebrow(s, 'Três pacotes. Acesso vitalício em todos.')
pac = [('pacote-hipocrates.jpg', 'Hipócrates', 'A base. Onde tudo começa.'),
       ('pacote-galeno.jpg', 'Galeno', 'O raciocínio clínico além da referência.'),
       ('pacote-olimpo.jpg', 'Olimpo', 'Os dois juntos, e a Avançada de volta.')]
for i, (f, n, d) in enumerate(pac):
    x = 0.6 + i * 4.1
    img(s, f, x, 1.3, w=3.85)
    txt(s, x, 4.05, 3.85, 1, [[(n, SERIF, 48, 'ouro2' if n != 'Olimpo' else 'rosa')]])
    txt(s, x, 5.0, 3.85, 1.2, [[(d, SANS, 18, 'texto')]], esp=1.1)
nota(s, '"Dei nome de quem olhou o paciente inteiro antes de todo mundo." Ainda sem preço.')


def pacote(nome, sub, itens, total, extra=None, foto=None, cor='ouro2', passo=0.75):
    s = novo('bg2'); eyebrow(s, sub)
    txt(s, 0.6, 1.0, 7.5, 1.3, [[(nome, SERIF, 72, cor)]])
    y = 2.45
    for t, v in itens:
        txt(s, 0.6, y, 6.2, 0.6, [[(t, SANS, 20, 'texto')]], anc='m')
        txt(s, 6.8, y, 1.8, 0.6, [[(v, SANS, 20, 'suave')]], al='r', anc='m')
        caixa(s, 0.6, y + passo - 0.09, 8.0, 0.01, cor='linha'); y += passo
    if extra: txt(s, 0.6, y + 0.1, 8.0, 0.8, [[(extra, SANS, 16, 'ouro', False, True)]], esp=1.1); y += 0.8
    txt(s, 0.6, 6.35, 5.5, 0.6, [[('Comprando separado', SANS, 18, 'suave')]], anc='m')
    txt(s, 5.6, 6.2, 3.0, 0.8, [[(total, SERIF, 44, 'creme')]], al='r', anc='m')
    if foto: img(s, foto, 9.0, 1.4, w=3.75)
    return s


pacote('Hipócrates', 'A base. Onde tudo começa.',
       [('Formação em Nutrição Funcional 3.0', 'R$ 3.497'), ('Formação em Suplementação Individualizada', 'R$ 4.497'),
        ('Jejum: Ferramenta para a Saúde', 'R$ 997'), ('Hemograma: O Ponto Cego', 'R$ 497')], 'R$ 9.488', foto='marco-hipocrates.jpg')
pacote('Galeno', 'O raciocínio clínico além da referência.',
       [('Formação em Exames Laboratoriais 2.0', 'R$ 3.497'), ('Formação em Saúde Hormonal', 'R$ 3.497'),
        ('Fertilidade e Gestação', 'R$ 1.997'), ('Alergias Alimentares', 'R$ 697')], 'R$ 9.688', foto='marco-galeno.jpg')
s = pacote('Olimpo', 'Hipócrates + Galeno, sem repetir aula, e a Avançada',
           [('Formação em Nutrição Funcional 3.0', 'R$ 3.497'), ('Formação em Suplementação Individualizada', 'R$ 4.497'),
            ('Jejum: Ferramenta para a Saúde', 'R$ 997'), ('Formação em Exames Laboratoriais 2.0', 'R$ 3.497'),
            ('Formação em Saúde Hormonal', 'R$ 3.497'), ('Fertilidade e Gestação', 'R$ 1.997'), ('Formação Avançada em Nutrição Funcional', 'R$ 1.988')], 'R$ 19.970', cor='rosa', passo=0.54)
caixa(s, 9.0, 1.3, 3.75, 4.6, cor='roxo', borda='ouro')
txt(s, 9.3, 1.55, 3.2, 4.2, [[('DE VOLTA, SÓ NO OLIMPO', SANS, 13, 'ouro', False, True)], [('Formação Avançada em Nutrição Funcional', SERIF, 30, 'creme')],
                            [('19 aulas · 72 h · com o Gabriel', SANS, 16, 'suave')], [('Valor de venda: R$ 1.988', SANS, 16, 'suave')],
                            [(' ', SANS, 8, 'suave')], [('Sem pagar duas vezes a mesma aula: Hemograma já está em Exames, e Alergias já está na NF 3.0.', SANS, 14, 'texto')]], esp=1.1)
nota(s, 'O total de R$ 19.970 inclui a Avançada (R$ 1.988). Mais de 300 horas de aula. Ainda sem o preço da Black.')

# 36 A CONTA DO VITALÍCIO
s = novo(); eyebrow(s, 'A conta do vitalício')
txt(s, 0.6, 1.0, 12, 1.1, [[('Para estudar 6 anos sem o vitalício, tu compraria a mesma formação 3 vezes.', SERIF, 36, 'creme')]], esp=1.0)
cab = ['Pacote', 'Hoje, 2 anos', 'Em 6 anos, comprando de novo']
lin = [('Hipócrates', 'R$ 9.488', 'R$ 28.464'), ('Galeno', 'R$ 9.688', 'R$ 29.064'), ('Olimpo', 'R$ 19.970', 'R$ 59.910')]
xs = [0.6, 4.6, 8.4]; ws = [3.8, 3.6, 4.3]
for j, c in enumerate(cab):
    txt(s, xs[j], 2.75, ws[j], 0.5, [[(c.upper(), SANS, 14, 'ouro', False, True)]], al='l' if j == 0 else 'r')
for i, l in enumerate(lin):
    y = 3.4 + i * 1.0
    caixa(s, 0.6, y + 0.9, 12.1, 0.01, cor='linha')
    for j, v in enumerate(l):
        txt(s, xs[j], y, ws[j], 0.85, [[(v, SERIF if j else SANS, 40 if j == 2 else 32, 'rosa' if j == 2 else ('texto' if j else 'creme'))]],
            al='l' if j == 0 else 'r', anc='m')
txt(s, 0.6, 6.6, 12, 0.5, [[('Simulação comprando de novo a cada 2 anos, pelos preços de hoje à vista.', SANS, 13, 'suave')]])
nota(s, 'Ainda sem o preço da Black. A Black entra só no último slide.')

# OBJEÇÕES
obj = [('"Gabriel, eu não sei se vou dar conta do conteúdo."', 'É gravado. Tu assiste no teu ritmo e volta quantas vezes precisar. Para sempre.'),
       ('"Gabriel, eu não tenho tempo agora."', 'É vitalício. Não tem relógio correndo. Tu começa quando puder.'),
       ('"Eu já tenho uma das formações."', 'Ela vem no pacote, agora vitalícia. E no Olimpo tu leva tudo o que ainda não tem, sem repetir aula.'),
       ('"Depois eu estudo."', 'Depois é o que a Narrativa Cega mais gosta de ouvir.'),
       ('"E se eu me arrepender?"', '15 dias de garantia incondicional. Assiste com calma. Se não for para ti, devolvemos 100% do valor.')]
for i, (q, r) in enumerate(obj):
    s = novo('bg2' if i % 2 else 'bg')
    eyebrow(s, 'Antes de decidir · %d/%d' % (i + 1, len(obj)))
    txt(s, 0.6, 1.4, 12.1, 2.4, [[(q, SERIF, 52, 'suave', True)]], anc='b', esp=1.0)
    caixa(s, 0.6, 4.15, 1.2, 0.04, cor='ouro')
    txt(s, 0.6, 4.5, 12.1, 2.3, [[(r, SERIF, 44, 'creme')]], esp=1.0)
    if i == 4: nota(s, 'A garantia fica antes do preço: tira o risco da decisão.')

# POR TI
s = novo(); frase(s, 'Faz por ti.', 'Por ti, que já viu a paciente voltar vezes demais.', sz=60)

# ÚLTIMO: PREÇO + CHAMADA
s = novo()
txt(s, 0.6, 0.4, 9, 0.5, [[('NA BLACK, PARA SEMPRE', SANS, 15, 'ouro', False, True)]])
col = [('Hipócrates', 'R$ 9.488', 'R$ 5.947', '12x R$ 579,40', 'ouro2'),
       ('Galeno', 'R$ 9.688', 'R$ 5.947', '12x R$ 579,40', 'ouro2'),
       ('Olimpo', 'R$ 19.970', 'R$ 9.947', '12x R$ 969,11', 'rosa')]
for i, (n, de, por, parc, cor) in enumerate(col):
    x = 0.6 + i * 3.05
    caixa(s, x, 1.05, 2.85, 4.4, cor='roxo' if n == 'Olimpo' else 'bg2', borda='rosa' if n == 'Olimpo' else 'linha')
    txt(s, x + 0.25, 1.25, 2.4, 0.8, [[(n, SERIF, 36, cor)]])
    t = txt(s, x + 0.25, 2.1, 2.4, 0.5, [[('separado ' + de, SANS, 15, 'suave')]])
    t.text_frame.paragraphs[0].runs[0].font._rPr.set('strike', 'sngStrike')
    txt(s, x + 0.25, 2.65, 2.4, 0.4, [[('à vista, com desconto', SANS, 13, 'suave')]])
    txt(s, x + 0.25, 3.0, 2.5, 0.9, [[(por, SERIF, 46, 'creme')]])
    txt(s, x + 0.25, 3.95, 2.4, 0.5, [[('ou ' + parc, SANS, 17, 'texto', False, True)]])
    if n == 'Olimpo': txt(s, x + 0.25, 4.5, 2.4, 0.8, [[('Um sexto do que custariam 6 anos. Com a Avançada.', SANS, 12, 'ouro')]], esp=1.05)
img(s, '/tmp/claude-0/-home-user/f6d0531c-9b3a-5a68-94b3-f4878f32a73b/scratchpad/slides/qr.png', 10.15, 1.05, w=2.55)
txt(s, 9.95, 3.75, 2.95, 1.0, [[('black.saudeavancada.com', SANS, 16, 'creme', False, True)]], al='c')
caixa(s, 0.6, 5.75, 12.1, 1.25, cor='rosa')
txt(s, 0.9, 5.75, 11.5, 1.25, [[('O vitalício fecha dia 31/10, às 23h59.  ', SANS, 24, 'creme', False, True),
                                ('Link aqui embaixo e na bio.  15 dias de garantia.', SANS, 18, 'creme')]], anc='m')
nota(s, 'ÚLTIMO SLIDE. Fica na tela durante todas as perguntas. Depois do preço não há mais conteúdo, só respostas e chamada para o link.')

prs.save('/tmp/claude-0/-home-user/f6d0531c-9b3a-5a68-94b3-f4878f32a73b/scratchpad/slides/Live-Black-FSA-19-10.pptx')
print(len(prs.slides), 'slides')
